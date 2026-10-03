#!/usr/bin/env python3
"""Export the Parthia HF Knowledge Base (Excel) to versioned, schema-validated JSON.

Offline authoring tool. The analytics runtime never reads the Excel file; it
reads only analytics/knowledge/json/*.json.

    python3 analytics/knowledge/export_knowledge.py            # export + validate
    python3 analytics/knowledge/export_knowledge.py --check    # fail if JSON is out of date

Rules:
- One JSON file per sheet. Every file carries the same knowledge_version
  (first 12 hex chars of the workbook's SHA-256), so a change to the Excel
  always produces a new version. No timestamps: the export is deterministic.
- Textual values are kept verbatim. Where a value can be parsed without
  interpretation (ID lists, "90 days", "< 3.5") a structured field is added
  next to it. Anything that needs clinical interpretation is NOT parsed here;
  it lives in analytics/knowledge/overrides.json, decided by the clinical lead.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path

import jsonschema
import openpyxl

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "source" / "Parthia Health - HF Knowledge Base.xlsx"
OUT = HERE / "json"
SCHEMAS = HERE / "schemas"

REVIEW_STATUS = {
    "draft – verify": "draft_verify",
    "draft - verify": "draft_verify",
    "draft": "draft_verify",
    "in review": "in_review",
    "approved": "approved",
    "retired": "retired",
}
SEVERITY = {"urgent": "urgent", "high": "high", "moderate": "moderate", "low": "low"}
ID_PATTERNS = {
    "element": r"\bD\d{3}\b",
    "rule": r"\bR\d{2}\b",
    "source": r"\bS\d{2}\b",
    "threshold": r"\bT\d{2}\b",
    "card": r"\bC\d{2}\b",
}
COMPARISON = re.compile(r"(≤|≥|<=|>=|<|>)\s*(\d+(?:\.\d+)?)")
OPS = {"≤": "<=", "≥": ">=", "<=": "<=", ">=": ">=", "<": "<", ">": ">"}

# Sheet name -> (output file stem, id column or None, column -> key mapping)
TABULAR = {
    "Data dictionary": ("data_dictionary", "Element ID", {
        "Element ID": "element_id", "Data element": "name", "Category": "category",
        "Definition": "definition", "Data type": "data_type", "Unit": "unit",
        "Code system": "code_system", "Code (verify)": "code", "Data source": "data_source",
        "Frequency": "frequency", "Valid for analysis (verify)": "valid_for_analysis",
        "Priority": "priority", "Required by rules": "required_by_rules",
        "Used by agent": "used_by_agent", "Justification (source ID)": "source_ids_text",
        "Status": "review_status_text", "Notes": "notes",
    }),
    "Thresholds": ("thresholds", "Threshold ID", {
        "Threshold ID": "threshold_id", "Element ID": "element_id", "Measure": "measure",
        "Unit": "unit", "Usual reference range": "reference_range",
        "Target in heart failure": "hf_target", "Alert LOW": "alert_low",
        "Alert HIGH": "alert_high", "URGENT (red flag)": "urgent",
        "Context / conditions": "context", "Patient-specific override?": "patient_override",
        "Source ID": "source_ids_text", "Status": "review_status_text", "Notes": "notes",
    }),
    "Rules": ("rules", "Rule ID", {
        "Rule ID": "rule_id", "Rule name": "name", "Category": "category",
        "Drugs / classes involved": "drugs", "Data needed": "data_needed",
        "Condition (logic)": "condition", "Threshold values": "threshold_values",
        "Severity": "severity_text", "Question for the patient to ask": "patient_question",
        "Note for the clinician": "clinician_note", "Source ID": "source_ids_text",
        "Section / page": "section", "Agent": "agent", "Applies to phenotype": "phenotype",
        "Data sources": "data_sources", "Comments": "comments",
        "Review status": "review_status_text", "Reviewed by": "reviewed_by",
        "Date reviewed": "date_reviewed", "Domain": "kb_domain",
    }),
    "Recommendation cards": ("recommendation_cards", "Card ID", {
        "Card ID": "card_id", "Topic": "topic", "Recommendation (own words)": "recommendation",
        "Class of recommendation": "class_of_recommendation",
        "Level of evidence": "level_of_evidence", "Population / phenotype": "population",
        "Drug / class tags": "drug_tags", "Lab / data tags": "data_tags",
        "Used by agent": "used_by_agent", "Patient-friendly version": "patient_text",
        "Source ID": "source_ids_text", "Section / page": "section",
        "Review status": "review_status_text", "Reviewed by": "reviewed_by",
        "Date reviewed": "date_reviewed", "Related rule IDs": "related_rules_text",
        "Domain": "kb_domain",
    }),
    "Source register": ("source_register", "Source ID", {
        "Source ID": "source_id", "Title": "title", "Organization": "organization",
        "Year / version": "version", "Source type": "source_type", "Scope": "scope",
        "Audience": "audience", "Link or DOI": "link", "License status": "license_status",
        "Full text storable": "full_text_storable", "Source status": "source_status",
        "Date accessed": "date_accessed", "Notes": "notes",
    }),
    "Test cases": ("test_cases", "Test ID", {
        "Test ID": "test_id", "Domain": "kb_domain", "Scenario (synthetic patient)": "scenario",
        "Input data": "input_data", "Expected rule IDs": "expected_rules_text",
        "Expected severity": "expected_severity_text",
        "Expected question / output": "expected_output", "Must NOT say": "must_not_say",
        "Source ID": "source_ids_text", "Status": "status", "Last result": "last_result",
        "Notes": "notes",
    }),
    "Domain tracker": ("domain_tracker", "Domain", {
        "Domain": "kb_domain", "Key question for the agents": "key_question", "Week": "week",
        "Main sources": "source_ids_text", "Target rules": "target_rules",
        "Rules approved": "rules_approved", "Target cards": "target_cards",
        "Cards approved": "cards_approved", "Target tests": "target_tests",
        "Tests ready or passed": "tests_ready", "Status": "status", "Notes": "notes",
    }),
}
ID_FORMATS = {
    "data_dictionary": r"^D\d{3}$", "thresholds": r"^T\d{2}$", "rules": r"^R\d{2}$",
    "recommendation_cards": r"^C\d{2}$", "source_register": r"^S\d{2}$",
    "test_cases": r"^TC\d{2}$", "domain_tracker": r"^\d+\. ",
}


def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    if hasattr(v, "isoformat"):
        return v.isoformat()[:10]
    return v


def ids(text, kind):
    return sorted(set(re.findall(ID_PATTERNS[kind], text or ""))) if text else []


def review_status(text):
    if text is None:
        return None
    return REVIEW_STATUS.get(text.strip().lower())


def parse_duration(text):
    """'90 days' / '12 months' / '7 days' -> {amount, unit}. Only the leading clause."""
    if not text:
        return None
    first = text.split(";")[0].strip()
    m = re.fullmatch(r"(\d+)\s*(day|days|month|months)", first, re.I)
    if not m:
        return None
    unit = "days" if m.group(2).lower().startswith("day") else "months"
    return {"amount": int(m.group(1)), "unit": unit}


def parse_comparisons(text):
    """Extract comparisons from a threshold cell.

    parse_status: "exact" (only comparisons joined by 'or'), "partial"
    (comparisons plus qualifying text: needs an override to be used),
    "none" (no comparison found).
    """
    if not text or text.strip().lower() in ("n/a", "na"):
        return {"comparisons": [], "parse_status": "not_applicable"}
    comps = [{"op": OPS[op], "value": float(v)} for op, v in COMPARISON.findall(text)]
    if not comps:
        return {"comparisons": [], "parse_status": "none"}
    residue = COMPARISON.sub("", text)
    residue = re.sub(r"\bor\b", "", residue).strip(" ,;")
    return {"comparisons": comps, "parse_status": "exact" if not residue else "partial"}


def read_table(ws, header_map, id_col):
    rows = list(ws.iter_rows(values_only=True))
    header = [clean(c) for c in rows[0]]
    missing = [h for h in header_map if h not in header]
    if missing:
        raise SystemExit(f"Sheet '{ws.title}': missing columns {missing}")
    out, notes = [], []
    for raw in rows[1:]:
        cells = dict(zip(header, (clean(c) for c in raw)))
        if all(v is None for v in cells.values()):
            continue
        rec = {key: cells.get(col) for col, key in header_map.items()}
        rid = rec[header_map[id_col]]
        fmt = ID_FORMATS[TABULAR[ws.title][0]]
        if rid is None or not re.match(fmt, str(rid)):
            # Free-text footnotes under the table (e.g. threshold priority order).
            notes.append(" ".join(str(v) for v in raw if v is not None))
            continue
        out.append(rec)
    return out, notes


def enrich(stem, rec):
    if "source_ids_text" in rec:
        rec["source_ids"] = ids(rec["source_ids_text"], "source")
    if "review_status_text" in rec:
        rec["review_status"] = review_status(rec["review_status_text"])
    if stem == "data_dictionary":
        rec["required_by_rule_ids"] = ids(rec["required_by_rules"], "rule")
        rec["validity"] = parse_duration(rec["valid_for_analysis"])
        rec["validity_has_context"] = bool(rec["valid_for_analysis"] and ";" in rec["valid_for_analysis"])
        rec["mvp"] = (rec["priority"] or "").lower() == "mvp"
    elif stem == "thresholds":
        for col in ("alert_low", "alert_high", "urgent"):
            rec[f"{col}_parsed"] = parse_comparisons(rec[col])
    elif stem == "rules":
        rec["severity"] = SEVERITY.get((rec["severity_text"] or "").lower())
        rec["element_ids"] = ids(" ".join(filter(None, [rec["data_needed"], rec["threshold_values"]])), "element")
        rec["threshold_ids"] = ids(rec["threshold_values"], "threshold")
    elif stem == "recommendation_cards":
        rec["related_rule_ids"] = ids(rec["related_rules_text"], "rule")
    elif stem == "test_cases":
        rec["expected_rule_ids"] = ids(rec["expected_rules_text"], "rule")
        rec["expected_severity"] = SEVERITY.get((rec["expected_severity_text"] or "").lower())
    return rec


def read_free_text(ws):
    return [[clean(c) for c in r if clean(c) is not None]
            for r in ws.iter_rows(values_only=True) if any(clean(c) is not None for c in r)]


def read_lists(ws):
    rows = list(ws.iter_rows(values_only=True))
    header = [clean(c) for c in rows[0]]
    lists = {h: [] for h in header if h}
    for raw in rows[1:]:
        for h, v in zip(header, raw):
            if h and clean(v) is not None:
                lists[h].append(clean(v))
    return lists


# ---- schemas ---------------------------------------------------------------

STR = {"type": ["string", "null"]}
RS = {"enum": ["draft_verify", "in_review", "approved", "retired", None]}
IDLIST = {"type": "array", "items": {"type": "string"}}


def row_schema(stem, keys):
    props = {k: {} for k in keys}
    required = list(keys)
    id_key = {"data_dictionary": "element_id", "thresholds": "threshold_id", "rules": "rule_id",
              "recommendation_cards": "card_id", "source_register": "source_id",
              "test_cases": "test_id", "domain_tracker": "kb_domain"}[stem]
    props[id_key] = {"type": "string", "pattern": ID_FORMATS[stem]}
    if "review_status" in keys:
        props["review_status"] = RS
    if "source_ids" in keys:
        props["source_ids"] = IDLIST
    if stem == "rules":
        props["severity"] = {"enum": ["urgent", "high", "moderate", "low"]}
        props["patient_question"] = {"type": "string"}
        props["review_status"] = {"enum": ["draft_verify", "in_review", "approved", "retired"]}
    if stem == "thresholds":
        props["element_id"] = {"type": "string", "pattern": r"^D\d{3}$"}
    if stem == "data_dictionary":
        props["review_status"] = {"enum": ["draft_verify", "in_review", "approved", "retired"]}
    return {"type": "object", "properties": props, "required": required, "additionalProperties": False}


def file_schema(stem, row):
    return {
        "$schema": "https://json-schema.org/draft/2020-12/schema",
        "$id": f"parthia/knowledge/{stem}.schema.json",
        "title": f"Parthia HF knowledge base: {stem}",
        "type": "object",
        "required": ["knowledge_version", "source_workbook", "sheet", "rows"],
        "additionalProperties": False,
        "properties": {
            "knowledge_version": {"type": "string", "pattern": "^[0-9a-f]{12}$"},
            "source_workbook": {"type": "string"},
            "sheet": {"type": "string"},
            "notes": {"type": "array", "items": {"type": "string"}},
            "rows": row,
        },
    }


def build():
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()[:12]
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    expected = set(TABULAR) | {"Dictionary guide", "Lists", "How to use"}
    if set(wb.sheetnames) != expected:
        raise SystemExit(f"Unexpected sheets. Found {wb.sheetnames}, expected {sorted(expected)}")
    files, schemas = {}, {}
    meta = {"knowledge_version": digest, "source_workbook": SOURCE.name}
    for sheet, (stem, id_col, header_map) in TABULAR.items():
        rows, notes = read_table(wb[sheet], header_map, id_col)
        rows = [enrich(stem, r) for r in rows]
        files[stem] = {**meta, "sheet": sheet, "notes": notes, "rows": rows}
        schemas[stem] = file_schema(stem, {"type": "array", "minItems": 1,
                                           "items": row_schema(stem, list(rows[0].keys()))})
    for sheet, stem in (("Dictionary guide", "dictionary_guide"), ("How to use", "how_to_use")):
        files[stem] = {**meta, "sheet": sheet, "rows": read_free_text(wb[sheet])}
        schemas[stem] = file_schema(stem, {"type": "array", "items": {"type": "array", "items": {"type": ["string", "number"]}}})
    files["lists"] = {**meta, "sheet": "Lists", "rows": read_lists(wb["Lists"])}
    schemas["lists"] = file_schema("lists", {"type": "object", "additionalProperties": IDLIST})
    files["manifest"] = {**meta, "sheet": "(manifest)", "rows": sorted(files)}
    schemas["manifest"] = file_schema("manifest", IDLIST)
    for stem, data in files.items():
        jsonschema.validate(data, schemas[stem])
    check_references(files)
    return files, schemas


def check_references(files):
    """Every cited ID must exist. Fails the export rather than shipping dangling references."""
    known = {
        "element": {r["element_id"] for r in files["data_dictionary"]["rows"]},
        "rule": {r["rule_id"] for r in files["rules"]["rows"]},
        "source": {r["source_id"] for r in files["source_register"]["rows"]},
        "threshold": {r["threshold_id"] for r in files["thresholds"]["rows"]},
    }
    errors = []
    for stem, data in files.items():
        if not isinstance(data["rows"], list):
            continue
        for row in data["rows"]:
            if not isinstance(row, dict):
                continue
            for key, kind in (("source_ids", "source"), ("element_ids", "element"),
                              ("threshold_ids", "threshold"), ("required_by_rule_ids", "rule"),
                              ("related_rule_ids", "rule"), ("expected_rule_ids", "rule")):
                for ref in row.get(key, []):
                    if ref not in known[kind]:
                        errors.append(f"{stem}: {ref} cited but not defined")
    if errors:
        raise SystemExit("Reference check failed:\n  " + "\n  ".join(errors))


def dump(obj):
    return json.dumps(obj, ensure_ascii=False, indent=2) + "\n"


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--check", action="store_true", help="fail if committed JSON differs from the workbook")
    args = ap.parse_args()
    files, schemas = build()
    outputs = {OUT / f"{s}.json": dump(d) for s, d in files.items()}
    outputs |= {SCHEMAS / f"{s}.schema.json": dump(d) for s, d in schemas.items()}
    if args.check:
        stale = [str(p.relative_to(HERE)) for p, text in outputs.items()
                 if not p.exists() or p.read_text(encoding="utf-8") != text]
        if stale:
            print("Knowledge JSON is out of date; run export_knowledge.py:\n  " + "\n  ".join(stale))
            return 1
        print(f"Knowledge JSON up to date (version {files['manifest']['knowledge_version']}).")
        return 0
    OUT.mkdir(exist_ok=True)
    SCHEMAS.mkdir(exist_ok=True)
    for p, text in outputs.items():
        p.write_text(text, encoding="utf-8")
    print(f"Exported {len(files)} files, knowledge_version {files['manifest']['knowledge_version']}.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
